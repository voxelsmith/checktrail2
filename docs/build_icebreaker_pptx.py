#!/usr/bin/env python3
"""Render Checktrail icebreaker slides in the live scribble look, then pack a PPTX."""

from __future__ import annotations

import math
import random
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches

W, H = 1920, 1080
BLACK = (0, 0, 0)
WHITE = (255, 255, 255)
MUTED = (200, 200, 200)
ROOT = Path(__file__).resolve().parent
FONT_DIR = ROOT / "fonts"
if not (FONT_DIR / "GochiHand-Regular.ttf").exists():
    FONT_DIR = Path("/tmp/checktrail-fonts")
OUT_DIR = Path("/tmp/checktrail-slides")
PPTX_PATH = ROOT / "Checktrail-Icebreaker.pptx"


def font(name: str, size: int) -> ImageFont.FreeTypeFont:
    file = "GochiHand-Regular.ttf" if name == "gochi" else "PatrickHand-Regular.ttf"
    return ImageFont.truetype(str(FONT_DIR / file), size)


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_w: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        trial = word if not line else f"{line} {word}"
        if draw.textlength(trial, font=fnt) <= max_w:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def scribble_rect(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], rng: random.Random, width: int = 3) -> None:
    x0, y0, x1, y1 = box
    jitter = 5

    def edge(x_a, y_a, x_b, y_b, steps: int) -> list[tuple[int, int]]:
        pts = []
        for i in range(steps + 1):
            t = i / steps
            x = x_a + (x_b - x_a) * t + rng.randint(-jitter, jitter)
            y = y_a + (y_b - y_a) * t + rng.randint(-jitter, jitter)
            pts.append((int(x), int(y)))
        return pts

    pts = []
    pts += edge(x0, y0, x1, y0, 18)
    pts += edge(x1, y0, x1, y1, 10)
    pts += edge(x1, y1, x0, y1, 18)
    pts += edge(x0, y1, x0, y0, 10)
    draw.line(pts + [pts[0]], fill=WHITE, width=width, joint="curve")


def text_block(
    draw: ImageDraw.ImageDraw,
    text: str,
    xy: tuple[int, int],
    fnt: ImageFont.FreeTypeFont,
    fill=WHITE,
    max_w: int | None = None,
    leading: float = 1.08,
) -> int:
    x, y = xy
    lines = wrap(draw, text, fnt, max_w) if max_w else text.split("\n")
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += int(fnt.size * leading)
    return y


def new_slide(seed: int) -> tuple[Image.Image, ImageDraw.ImageDraw, random.Random]:
    rng = random.Random(seed)
    img = Image.new("RGB", (W, H), BLACK)
    noise = Image.effect_noise((W, H), 14).convert("L")
    grain = Image.merge("RGB", (noise, noise, noise))
    img = Image.blend(img, grain, 0.05)
    draw = ImageDraw.Draw(img)
    # tiny margin frame, wobbly
    scribble_rect(draw, (28, 24, W - 28, H - 24), rng, width=2)
    return img, draw, rng


def footer(draw: ImageDraw.ImageDraw, page: int, total: int) -> None:
    fnt = font("patrick", 28)
    draw.text((64, H - 70), "Checktrail", font=fnt, fill=MUTED)
    label = f"{page} / {total}"
    w = draw.textlength(label, font=fnt)
    draw.text((W - 64 - w, H - 70), label, font=fnt, fill=MUTED)


def measure(draw, text, fnt, max_w, leading: float) -> int:
    lines = wrap(draw, text, fnt, max_w)
    return int(len(lines) * fnt.size * leading)


def card(draw, rng, box, title, body, title_size=42, body_size=32) -> None:
    scribble_rect(draw, box, rng, width=3)
    x0, y0, x1, y1 = box
    inner = x1 - x0 - 72
    ty = text_block(draw, title, (x0 + 36, y0 + 28), font("gochi", title_size), WHITE, inner, 1.05)
    text_block(draw, body, (x0 + 36, ty + 10), font("patrick", body_size), MUTED, inner, 1.12)


def fitted_row(draw, rng, y, items, title_size=42, body_size=32, gap=40, x0=90, x1=1830, min_h=280) -> None:
    n = len(items)
    usable = x1 - x0 - gap * (n - 1)
    w = usable // n
    heights = []
    for title, body in items:
        inner = w - 72
        h = 28 + measure(draw, title, font("gochi", title_size), inner, 1.05)
        h += 10 + measure(draw, body, font("patrick", body_size), inner, 1.12) + 36
        heights.append(max(min_h, h))
    h = max(heights)
    for i, (title, body) in enumerate(items):
        left = x0 + i * (w + gap)
        card(draw, rng, (left, y, left + w, y + h), title, body, title_size, body_size)


def save(img: Image.Image, name: str) -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{name}.png"
    img.save(path, "PNG", optimize=True)
    return path


def build() -> Path:
    paths: list[Path] = []
    total = 10

    # 1 cover
    img, draw, rng = new_slide(1)
    draw.text((90, 150), "Checktrail", font=font("gochi", 132), fill=WHITE)
    draw.text((96, 300), "Party games that unfreeze a room.", font=font("gochi", 64), fill=WHITE)
    scribble_rect(draw, (90, 430, 980, 760), rng)
    text_block(
        draw,
        "Two phone games. One invite link. The questions do the introducing so nobody has to.",
        (124, 470),
        font("patrick", 42),
        MUTED,
        800,
        1.15,
    )
    scribble_rect(draw, (1080, 430, 1780, 760), rng)
    text_block(draw, "Anon Wheel", (1120, 470), font("gochi", 48), WHITE, 600)
    text_block(draw, "Secret questions. A spinning wheel. A buzzer finale.", (1120, 540), font("patrick", 36), MUTED, 600)
    text_block(draw, "Mirror Vote", (1120, 640), font("gochi", 48), WHITE, 600)
    text_block(draw, "Anonymous “most likely” votes. A portrait of how the room sees you.", (1120, 710), font("patrick", 32), MUTED, 600)
    footer(draw, 1, total)
    paths.append(save(img, "01_cover"))

    # 2 the problem
    img, draw, rng = new_slide(2)
    draw.text((90, 80), "Icebreakers usually freeze people.", font=font("gochi", 64), fill=WHITE)
    text_block(
        draw,
        "A circle of strangers. A host with a script. Names, jobs, “fun facts.” Everyone waits for it to be over.",
        (90, 180),
        font("patrick", 38),
        MUTED,
        1700,
        1.15,
    )
    fitted_row(
        draw,
        rng,
        340,
        [
            ("Too much spotlight", "Being asked to introduce yourself is a performance. Shy people go quiet. Loud people fill the air."),
            ("Too much honesty, too soon", "People protect themselves. They give safe answers. The room learns nothing real."),
            ("No shared dare", "Without a game, there is no reason to ask the interesting question — or to answer it."),
        ],
        42,
        32,
        min_h=300,
    )
    footer(draw, 2, total)
    paths.append(save(img, "02_problem"))

    # 3 the idea
    img, draw, rng = new_slide(3)
    draw.text((90, 80), "The idea", font=font("gochi", 72), fill=WHITE)
    text_block(
        draw,
        "Checktrail turns a pile of phones into a party. The host picks a game, creates a room, and sends a link. Friends join that room. From then on the rules — not the host — decide who speaks.",
        (90, 190),
        font("patrick", 40),
        MUTED,
        1740,
        1.14,
    )
    fitted_row(
        draw,
        rng,
        430,
        [
            ("Low friction", "No accounts. A name, a room, a link. People are playing in under a minute."),
            ("Shared dare", "A wheel, a vote, a buzzer. Structure that makes asking and answering feel like a game, not an interview."),
            ("Safety first", "Questions and votes can stay anonymous until the room is ready. Honesty without a name on it."),
        ],
        42,
        32,
        min_h=280,
    )
    footer(draw, 3, total)
    paths.append(save(img, "03_idea"))

    # 4 how a night starts
    img, draw, rng = new_slide(4)
    draw.text((90, 80), "How a night starts", font=font("gochi", 70), fill=WHITE)
    steps = [
        ("1", "Pick a category", "The host chooses Anon Wheel or Mirror Vote. That choice sets the mood: loud and messy, or quiet and revealing."),
        ("2", "Make a room", "One person creates. Everyone else taps the invite. Names appear in the lobby. The room is already a group."),
        ("3", "Play on every phone", "Same live room. Same round. No passing one device around the table. Nobody is left watching."),
        ("4", "Let the game talk", "The host is not the icebreaker. The questions, the wheel, and the votes are."),
    ]
    for i, (num, title, body) in enumerate(steps):
        y = 210 + i * 180
        scribble_rect(draw, (90, y, 1830, y + 155), rng)
        draw.text((120, y + 28), num, font=font("gochi", 64), fill=WHITE)
        draw.text((220, y + 22), title, font=font("gochi", 44), fill=WHITE)
        text_block(draw, body, (220, y + 80), font("patrick", 32), MUTED, 1540, 1.1)
    footer(draw, 4, total)
    paths.append(save(img, "04_night"))

    # 5 two games
    img, draw, rng = new_slide(5)
    draw.text((90, 80), "Two games. Two kinds of courage.", font=font("gochi", 64), fill=WHITE)
    scribble_rect(draw, (90, 220, 920, 780), rng)
    text_block(draw, "CATEGORY 1", (130, 260), font("patrick", 28), MUTED)
    text_block(draw, "Anon Wheel", (130, 310), font("gochi", 72), WHITE)
    text_block(
        draw,
        "The room writes questions in secret. A wheel picks who must answer. Skip too many times and you’re out. The last two smash a giant buzzer in rapid fire. At the end, you find out who asked what.",
        (130, 420),
        font("patrick", 36),
        MUTED,
        740,
        1.14,
    )
    scribble_rect(draw, (1000, 220, 1830, 780), rng)
    text_block(draw, "CATEGORY 2", (1040, 280), font("patrick", 28), MUTED)
    text_block(draw, "Mirror Vote", (1040, 330), font("gochi", 72), WHITE)
    text_block(
        draw,
        "Everyone gets the same “most likely” questions, shuffled. Votes stay anonymous. Then the room sees charts — and a trait portrait of how people built you.",
        (1040, 440),
        font("patrick", 36),
        MUTED,
        740,
        1.14,
    )
    footer(draw, 5, total)
    paths.append(save(img, "05_two_games"))

    # 6 Anon Wheel icebreak
    img, draw, rng = new_slide(6)
    draw.text((90, 70), "Anon Wheel breaks the ice by hiding the asker.", font=font("gochi", 56), fill=WHITE)
    points = [
        ("Secret questions", "People ask what they actually want to know — crush stuff, weird habits, the story behind the quiet person — because their name is not on the card yet."),
        ("The wheel is the host", "Nobody has to pick on a friend. Chance does it. That keeps the dare fair, and it keeps shy players from being hunted."),
        ("Skips have a cost", "Three skips and you’re out. You can dodge once. You cannot hide all night. The game gently forces people into the circle."),
        ("Rapid fire, then reveal", "The finale is loud and competitive. The author reveal comes last — after the room has already laughed. Identity is a punchline, not a threat."),
    ]
    fitted_row(draw, rng, 200, points[:2], 46, 34, min_h=250)
    fitted_row(draw, rng, 520, points[2:], 46, 34, min_h=250)
    footer(draw, 6, total)
    paths.append(save(img, "06_anon_ice"))

    # 7 Mirror Vote icebreak
    img, draw, rng = new_slide(7)
    draw.text((90, 70), "Mirror Vote breaks the ice by showing you to yourself.", font=font("gochi", 52), fill=WHITE)
    points = [
        ("Anonymous “most likely”", "Classic party questions, without the raised hands. You can name someone without owning the vote in front of them."),
        ("Shuffled decks", "Everyone answers the same prompts in a different order. You cannot copy the person next to you. The portrait is yours."),
        ("The room’s opinion, graphed", "Charts turn gossip into a shared object. People talk about the picture, not about confronting each other."),
        ("A trait portrait", "You leave with a sketch of how the group built you. That is a conversation that continues after the phones go down."),
    ]
    fitted_row(draw, rng, 200, points[:2], 46, 34, min_h=250)
    fitted_row(draw, rng, 520, points[2:], 46, 34, min_h=250)
    footer(draw, 7, total)
    paths.append(save(img, "07_mirror_ice"))

    # 8 social design
    img, draw, rng = new_slide(8)
    draw.text((90, 80), "What actually makes strangers talk", font=font("gochi", 62), fill=WHITE)
    rows = [
        ("Anonymity with a timer", "Honesty first. Names later. People risk a real question because the social cost is delayed."),
        ("A shared object", "A wheel, a chart, a buzzer. Attention goes to the game, not to “please like me.”"),
        ("Equal phones", "Every player has the same controls. No one is stuck as the audience. That is how mixed groups (shy + loud, new + old friends) stay in one conversation."),
        ("Permission to be uncool", "The scribble look is a notebook, not a product demo. It tells the room: this is play, not a networking event."),
    ]
    for i, (title, body) in enumerate(rows):
        y = 200 + i * 185
        scribble_rect(draw, (90, y, 1830, y + 160), rng)
        draw.text((130, y + 18), title, font=font("gochi", 42), fill=WHITE)
        text_block(draw, body, (130, y + 78), font("patrick", 32), MUTED, 1620, 1.1)
    footer(draw, 8, total)
    paths.append(save(img, "08_social"))

    # 9 who it's for
    img, draw, rng = new_slide(9)
    draw.text((90, 80), "Built for rooms that don’t know how to start", font=font("gochi", 56), fill=WHITE)
    fitted_row(
        draw,
        rng,
        220,
        [
            ("New groups", "First-week classmates. Game-jam teams. A dinner where half the table just met. Checktrail gives them a script that does not feel like a script."),
            ("Mixed energy", "The loud ones get a wheel and a buzzer. The quiet ones get anonymous questions and secret votes. Both stay in the same room."),
            ("After the phones", "The point is not the app. The point is the story you can retell: “the wheel landed on you,” “the room drew me as chaotic.” That is the ice, broken."),
        ],
        42,
        32,
        min_h=360,
    )
    footer(draw, 9, total)
    paths.append(save(img, "09_who"))

    # 10 close
    img, draw, rng = new_slide(10)
    draw.text((90, 160), "Don’t introduce yourselves.", font=font("gochi", 72), fill=WHITE)
    draw.text((90, 260), "Play.", font=font("gochi", 120), fill=WHITE)
    scribble_rect(draw, (90, 470, 1200, 820), rng)
    text_block(draw, "Host picks a game. Friends join the room. The questions do the rest.", (130, 520), font("patrick", 42), MUTED, 1000, 1.15)
    text_block(draw, "checktrail.vercel.app", (130, 700), font("gochi", 52), WHITE, 1000)
    scribble_rect(draw, (1280, 470, 1830, 820), rng)
    text_block(draw, "Anon Wheel", (1320, 520), font("gochi", 40), WHITE)
    text_block(draw, "Mirror Vote", (1320, 590), font("gochi", 40), WHITE)
    text_block(draw, "One link. Whole table.", (1320, 700), font("patrick", 34), MUTED, 450)
    footer(draw, 10, total)
    paths.append(save(img, "10_close"))

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in paths:
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0, 0, 0)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), prs.slide_width, prs.slide_height)

    PPTX_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(PPTX_PATH))
    return PPTX_PATH


if __name__ == "__main__":
    out = build()
    print(out)
